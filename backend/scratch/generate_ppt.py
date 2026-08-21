import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_bugflow_presentation(output_path):
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme colors
    BG_DARK = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)  # Slate 700
    PRIMARY = RGBColor(99, 102, 241)    # Indigo 500
    SECONDARY = RGBColor(6, 182, 212)   # Cyan 500
    TEXT_MAIN = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    ACCENT_GREEN = RGBColor(16, 185, 129)# Emerald 500

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background() # No border
        return bg

    def add_header(slide, title_text, category_text="BUGFLOW PROJECT PRESENTATION"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SECONDARY
        
        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    def add_card(slide, left, top, width, height, title, points, accent_color=PRIMARY):
        # Card Background Shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)
        
        # Title Box inside Card
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = accent_color
        p0.space_after = Pt(12)
        
        for pt in points:
            p = tf.add_paragraph()
            p.text = f"•  {pt}"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(8)

    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: Title Slide ====================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_background(slide1)
    
    # Title Box
    t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(3.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "BUGFLOW"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY
    p1.space_after = Pt(6)
    
    p2 = tf1.add_paragraph()
    p2.text = "Intelligent AI-Powered Defect Management & Vision Engine"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MAIN
    p2.space_after = Pt(16)
    
    p3 = tf1.add_paragraph()
    p3.text = "Automated Screenshot OCR Extraction • Gemini 2.5 AI Triage • Role-Based Engineering Workspaces"
    p3.font.size = Pt(14)
    p3.font.color.rgb = SECONDARY
    p3.space_after = Pt(28)
    
    p4 = tf1.add_paragraph()
    p4.text = "Software Engineering Senior Capstone / Final Project Presentation"
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 2: Problem Statement ====================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_background(slide2)
    add_header(slide2, "Problem Statement: Challenges of Legacy Defect Management")
    
    add_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "1. High Data Entry Overhead", 
             ["Legacy tools (Jira, Bugzilla) require manual typing of 10+ required fields.",
              "QA testers spend up to 40% of test cycles writing issue descriptions.",
              "Critical steps to reproduce are frequently omitted by mistake."])
              
    add_card(slide2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "2. Unstructured Context & Slow Triage", 
             ["Raw screenshots uploaded without text OCR parsing require manual inspection.",
              "Project leads spend hours matching bugs to developer skillsets.",
              "Duplicate bug reports cause backlog clutter and wasted engineering effort."])

    # ==================== SLIDE 3: Project Objectives ====================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_background(slide3)
    add_header(slide3, "Project Objectives & Core Deliverables")
    
    col_w = Inches(3.6)
    gap = Inches(0.3)
    add_card(slide3, Inches(0.8), Inches(1.8), col_w, Inches(5.0), 
             "⚡ Automated OCR Extraction", 
             ["Parse UI screenshots automatically via Gemini 2.5 Vision.",
              "Extract title, steps to reproduce, expected vs actual behavior.",
              "Eliminate manual description writing entirely."])

    add_card(slide3, Inches(0.8) + col_w + gap, Inches(1.8), col_w, Inches(5.0), 
             "🎯 Predictive AI Triage", 
             ["Auto-predict defect severity (Critical, High, Medium, Low).",
              "Suggest optimal developer assignment based on technical skills.",
              "Detect vector similarity to eliminate duplicate bugs."])

    add_card(slide3, Inches(0.8) + (col_w + gap)*2, Inches(1.8), col_w, Inches(5.0), 
             "🛡️ Enterprise RBAC & Design", 
             ["Tailored workspaces for Admin, Dev, Tester, and PM roles.",
              "User-controlled customizable theme engine.",
              "100% clean SaaS boot with real-time analytics metrics."])

    # ==================== SLIDE 4: Comparative Analysis ====================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_background(slide4)
    add_header(slide4, "Comparative Analysis: Legacy Systems vs. BugFlow AI")

    # Table creation
    rows, cols = 6, 3
    left, top, width, height = Inches(0.8), Inches(1.8), Inches(11.73), Inches(5.0)
    table_shape = slide4.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(4.2)
    table.columns[2].width = Inches(4.33)
    
    headers = ["Evaluation Metric", "Legacy Tools (Jira / Bugzilla)", "Proposed BugFlow Platform"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_MAIN
        
    matrix_data = [
        ["Bug Creation Speed", "Manual input (5–10 mins per bug)", "1-Click AI Screenshot OCR (< 5 seconds)"],
        ["Severity Estimation", "Subjective human manual picking", "Gemini Predictive AI Auto-Classification"],
        ["Developer Triage", "Manual dispatch by Project Manager", "AI Skill-Based Match Engine"],
        ["Resolution Assistance", "None (Dev investigates from scratch)", "AI Co-Pilot Root Cause & Fix Suggestions"],
        ["Data Boot Policy", "Requires fabricated demo data", "Clean Zero-Data Startup + Real Metrics"]
    ]
    
    for row_idx, row_data in enumerate(matrix_data, start=1):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = cell_value
            p.font.size = Pt(12)
            p.font.color.rgb = SECONDARY if col_idx == 2 else TEXT_MUTED

    # ==================== SLIDE 5: System Architecture ====================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_background(slide5)
    add_header(slide5, "System Architecture: Widescreen Tiered Overview")

    tiers = [
        ("1. Presentation Tier (Frontend)", "React 18 SPA • Vite Build • Tailwind CSS v4 • Lucide Icons • User-Controlled Theme Provider", SECONDARY),
        ("2. Application API Tier (Backend)", "Python 3.14 • Flask Modular Blueprints • Flask-JWT-Extended Security • REST API Contracts", PRIMARY),
        ("3. Intelligence Tier (AI Engine)", "Google Gemini 2.5 Flash SDK • Multimodal OCR Vision Parser • Code Fix Co-Pilot Engine", ACCENT_GREEN),
        ("4. Persistence Tier (Database)", "SQLAlchemy 2.0 ORM • SQLite / PostgreSQL Engine • Automated Relational Cascade Handlers", SECONDARY)
    ]
    
    y = Inches(1.8)
    for title, desc, col in tiers:
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.73), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)
        
        tb = slide5.shapes.add_textbox(Inches(1.1), y + Inches(0.15), Inches(11.1), Inches(0.8))
        tf = tb.text_frame
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = col
        
        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(13)
        p1.font.color.rgb = TEXT_MAIN
        
        y += Inches(1.3)

    # ==================== SLIDE 6: Core Platform Modules ====================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_background(slide6)
    add_header(slide6, "Core Platform Modules & Role-Based Workspaces")

    w = Inches(5.6)
    h = Inches(2.4)
    add_card(slide6, Inches(0.8), Inches(1.8), w, h, "🛡️ Admin Control Module", 
             ["Full system user management & role policies.",
              "Project creation, member allocation & key generation.",
              "Global security parameters & system configurations."], SECONDARY)

    add_card(slide6, Inches(6.8), Inches(1.8), w, h, "💻 Developer Workspace", 
             ["Personalized assigned bug queues & interactive Kanban board.",
              "AI Solution Co-Pilot providing root causes & code fixes.",
              "State transition triggers (In Progress -> Resolved)."], PRIMARY)

    add_card(slide6, Inches(0.8), Inches(4.5), w, h, "🧪 Tester Quality Module", 
             ["Multimodal screenshot bug extraction modal.",
              "AI description preview & manual override before saving.",
              "Bug verification & reopening/closure controls."], ACCENT_GREEN)

    add_card(slide6, Inches(6.8), Inches(4.5), w, h, "📊 Project Manager Analytics", 
             ["Real-time sprint velocity & health monitoring.",
              "Developer workload distribution & defect age reports.",
              "Release readiness summaries & analytics exports."], SECONDARY)

    # ==================== SLIDE 7: Screenshot Bug Extraction ====================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_background(slide7)
    add_header(slide7, "Multimodal Screenshot Defect OCR Engine Workflow")

    steps = [
        ("Step 1: Upload", "Tester uploads PNG/JPG UI screenshot"),
        ("Step 2: Payload", "Frontend packages base64 image data"),
        ("Step 3: Vision OCR", "Gemini 2.5 Flash parses visual pixel text"),
        ("Step 4: AI Extraction", "Extracts title, steps, expected vs actual"),
        ("Step 5: Triage", "Predicts severity & suggests developer match"),
        ("Step 6: Commit", "User reviews preview & commits to Database")
    ]
    
    col_w7 = Inches(3.6)
    gap7 = Inches(0.4)
    
    for idx, (st, desc) in enumerate(steps[:3]):
        add_card(slide7, Inches(0.8) + idx*(col_w7+gap7), Inches(1.8), col_w7, Inches(2.3), st, [desc], PRIMARY)
        
    for idx, (st, desc) in enumerate(steps[3:]):
        add_card(slide7, Inches(0.8) + idx*(col_w7+gap7), Inches(4.5), col_w7, Inches(2.3), st, [desc], SECONDARY)

    # ==================== SLIDE 8: Defect Lifecycle Workflow ====================
    slide8 = prs.slides.add_slide(blank_layout)
    apply_background(slide8)
    add_header(slide8, "Defect Lifecycle & Rigid State Machine Rules")

    states = [
        ("Reported", "Bug submitted by Tester; awaiting triage.", PRIMARY),
        ("In Review", "Evaluated by PM/Lead for validity.", SECONDARY),
        ("In Progress", "Assigned to Dev; actively under fix.", ACCENT_GREEN),
        ("Resolved", "Dev submitted code fix & resolution log.", SECONDARY),
        ("Closed", "Verified by Tester in staging environment.", PRIMARY)
    ]
    
    state_w = Inches(2.1)
    state_gap = Inches(0.3)
    x = Inches(0.8)
    
    for name, desc, col in states:
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.5), state_w, Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(2)
        
        tb = slide8.shapes.add_textbox(x + Inches(0.1), Inches(2.7), state_w - Inches(0.2), Inches(2.8))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = name
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = col
        p0.alignment = PP_ALIGN.CENTER
        p0.space_after = Pt(12)
        
        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEXT_MUTED
        p1.alignment = PP_ALIGN.CENTER
        
        x += state_w + state_gap

    # ==================== SLIDE 9: Technology Stack ====================
    slide9 = prs.slides.add_slide(blank_layout)
    apply_background(slide9)
    add_header(slide9, "Technology Stack & Software Engineering Tools")

    add_card(slide9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "🎨 Frontend Stack", 
             ["React 18 SPA Architecture",
              "Vite 8 Lightning Build Engine",
              "Tailwind CSS v4 Modern Styling",
              "Lucide React Icon Library",
              "Axios HTTP Interceptors"], SECONDARY)

    add_card(slide9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "⚙️ Backend Stack", 
             ["Python 3.14 Engine",
              "Flask 3.0 Web Framework",
              "SQLAlchemy 2.0 ORM",
              "Flask-JWT-Extended Authentication",
              "Werkzeug Password Hashing"], PRIMARY)

    add_card(slide9, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0), 
             "🤖 AI & Test Suite", 
             ["Google Gemini 2.5 Flash SDK",
              "Multimodal Vision API",
              "Pytest Test Suite (100% Pass Rate)",
              "In-Memory Isolated Testing (:memory:)",
              "GitHub Actions CI/CD Pipeline"], ACCENT_GREEN)

    # ==================== SLIDE 10: Results & Future Scope ====================
    slide10 = prs.slides.add_slide(blank_layout)
    apply_background(slide10)
    add_header(slide10, "Verified Results & Future Engineering Scope")

    add_card(slide10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "✅ Verified Project Results", 
             ["Zero-Data Boot: Clean database startup with 0 demo clutter.",
              "Fast Compilation: Vite production client build completed in 6.66s.",
              "100% Test Pass Rate: 6/6 Pytest suite tests passing in 8s.",
              "Theme Support: 6 user-controlled custom visual themes.",
              "Robust Security: Password hashing + JWT session persistence."], ACCENT_GREEN)

    add_card(slide10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "🚀 Future Engineering Roadmap", 
             ["Phase 1: Bi-directional GitHub & GitLab webhook integration.",
              "Phase 2: Automated AI Pull Request generation from bug fixes.",
              "Phase 3: CI/CD test failure auto-triage (Jenkins & GitHub Actions).",
              "Phase 4: Real-time WebSockets team chat & live collaboration."], SECONDARY)

    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.dirname(__file__))
    out_file = os.path.join(out_dir, "BugFlow_Presentation.pptx")
    create_bugflow_presentation(out_file)
